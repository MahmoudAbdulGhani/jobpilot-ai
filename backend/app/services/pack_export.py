"""Plain-text approved-document exports, isolated from providers and source files.

The worker accepts only validated block text, never paths, URLs, HTML or images.
Each export has a hard process deadline in addition to content/layout limits.
"""
from __future__ import annotations

import io
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
import time
from typing import Literal
import unicodedata
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
import reportlab
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate

MAX_BLOCKS = 100
MAX_BLOCK_CHARS = 2_000
MAX_DOCUMENT_CHARS = 30_000
MAX_PAGES = 20
MAX_EXPORT_BYTES = 5 * 1024 * 1024
EXPORT_TIMEOUT_SECONDS = 10
_WORKERS = threading.BoundedSemaphore(4)


class ExportFailure(Exception):
    def __init__(self, code: str, message: str | None = None):
        self.code = code
        self.message = message or str(code)
        super().__init__(self.message)


def _validated_blocks(document: dict) -> list[dict[str, str]]:
    blocks = document.get("blocks") if isinstance(document, dict) else None
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= MAX_BLOCKS:
        raise ExportFailure(
            "limit", f"Documents must contain 1–{MAX_BLOCKS} blocks.")
    result = []
    total = 0
    for block in blocks:
        if not isinstance(block, dict) or block.get("kind") not in {"heading", "paragraph", "bullet"}:
            raise ExportFailure(
                "invalid", "The approved document contains an invalid block.")
        value = block.get("text")
        if not isinstance(value, str) or not value.strip() or len(value) > MAX_BLOCK_CHARS:
            raise ExportFailure(
                "limit", "An approved document block is empty or exceeds 2,000 characters.")
        if any(unicodedata.category(char) in {"Cc", "Cs"} and char not in "\n\t" for char in value):
            raise ExportFailure(
                "invalid", "The document contains unsupported control characters.")
        total += len(value)
        if total > MAX_DOCUMENT_CHARS:
            raise ExportFailure(
                "limit", "The document exceeds the 30,000-character export limit.")
        # Deliberately drop identifiers, evidence, authorship and every audit field.
        result.append({"kind": block["kind"], "text": value})
    return result


def _check_deadline(deadline: float) -> None:
    if time.monotonic() > deadline:
        raise ExportFailure(
            "timeout", "Document export took too long. Try a shorter draft.")


class _BoundedBuffer(io.BytesIO):
    def write(self, data: bytes) -> int:
        if self.tell() + len(data) > MAX_EXPORT_BYTES:
            raise ExportFailure(
                "limit", "The exported document exceeds the 5 MB limit.")
        return super().write(data)


def _register_fonts() -> None:
    # Bitstream Vera is included, with its license, in the ReportLab package.
    font_dir = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("PackVera", "Vera.ttf"), ("PackVeraBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))


def _pdf(blocks: list[dict[str, str]], deadline: float) -> bytes:
    _register_fonts()
    styles = {
        "heading": ParagraphStyle(
            "Pack heading", fontName="PackVeraBold", fontSize=13, leading=17,
            textColor=HexColor("#183d31"), spaceBefore=12, spaceAfter=7,
            keepWithNext=True, allowWidows=0, allowOrphans=0,
        ),
        "paragraph": ParagraphStyle(
            "Pack body", fontName="PackVera", fontSize=10.5, leading=15,
            textColor=HexColor("#202c26"), spaceAfter=8,
            allowWidows=0, allowOrphans=0,
        ),
        "bullet": ParagraphStyle(
            "Pack bullet", fontName="PackVera", fontSize=10.5, leading=15,
            textColor=HexColor("#202c26"), spaceAfter=5, leftIndent=12,
            firstLineIndent=0, bulletIndent=0, bulletFontName="PackVera",
            allowWidows=0, allowOrphans=0,
        ),
    }
    story = []
    for block in blocks:
        _check_deadline(deadline)
        style = styles[block["kind"]]
        cmap = pdfmetrics.getFont(style.fontName).face.charToGlyph
        for char in block["text"]:
            if char not in "\n\t" and ord(char) not in cmap:
                raise ExportFailure(
                    "unsupported_glyph",
                    "The PDF font cannot display one or more characters in this draft. "
                    "Download DOCX, which preserves Unicode text, or edit those characters.",
                )
        # Paragraph supports markup, so all user text must be escaped first.
        value = escape(block["text"]).replace(
            "\n", "<br/>").replace("\t", "    ")
        story.append(
            Paragraph(value, style, bulletText="•" if block["kind"] == "bullet" else None))

    output = _BoundedBuffer()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=54, leftMargin=54,
        topMargin=46, bottomMargin=48, title="", author="", subject="",
        creator="", keywords="", pageCompression=1,
    )

    def page_footer(canvas, doc):
        _check_deadline(deadline)
        if doc.page > MAX_PAGES:
            raise ExportFailure(
                "limit", f"Exports are limited to {MAX_PAGES} pages.")
        canvas.saveState()
        canvas.setFont("PackVera", 8)
        canvas.setFillColor(HexColor("#637067"))
        canvas.drawRightString(A4[0] - 54, 27, str(doc.page))
        canvas.restoreState()

    try:
        document.build(story, onFirstPage=page_footer,
                       onLaterPages=page_footer)
    except ExportFailure:
        raise
    except Exception as error:
        raise ExportFailure(
            "render_failed", "The approved document could not be exported. Please retry.") from error
    _check_deadline(deadline)
    return output.getvalue()


def _docx(blocks: list[dict[str, str]], deadline: float) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(17), Mm(17)
    section.left_margin, section.right_margin = Mm(19), Mm(19)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("202C26")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    heading = document.styles["Heading 1"]
    heading.font.name, heading.font.size = "Arial", Pt(13)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string("183D31")
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)
    for block in blocks:
        _check_deadline(deadline)
        style = {"heading": "Heading 1", "paragraph": "Normal",
                 "bullet": "List Bullet"}[block["kind"]]
        paragraph = document.add_paragraph(block["text"], style)
        paragraph.paragraph_format.widow_control = True
        if block["kind"] == "bullet":
            paragraph.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = 2  # right aligned; only a page number, never internal metadata
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    properties = document.core_properties
    for attribute in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category", "identifier"):
        setattr(properties, attribute, "")
    output = _BoundedBuffer()
    document.save(output)
    _check_deadline(deadline)
    return output.getvalue()


def _render(document: dict, format: str) -> bytes:
    blocks = _validated_blocks(document)
    deadline = time.monotonic() + EXPORT_TIMEOUT_SECONDS
    if format == "pdf":
        return _pdf(blocks, deadline)
    if format == "docx":
        return _docx(blocks, deadline)
    raise ExportFailure("unsupported", "Choose PDF or DOCX for this download.")


def render_document(document: dict, format: Literal["pdf", "docx"]) -> bytes:
    """Render exact approved block text in a killable, resource-bounded worker."""
    blocks = _validated_blocks(document)
    if format not in {"pdf", "docx"}:
        raise ExportFailure(
            "unsupported", "Choose PDF or DOCX for this download.")
    if not _WORKERS.acquire(blocking=False):
        raise ExportFailure(
            "busy", "Other document downloads are being prepared. Please retry shortly.")
    try:
        process = subprocess.Popen(
            [sys.executable, "-m", "app.services.pack_export"],
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            cwd=str(Path(__file__).resolve().parents[2]),
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        try:
            output, error = process.communicate(
                json.dumps({"format": format, "document": {
                           "blocks": blocks}}).encode("utf-8"),
                timeout=EXPORT_TIMEOUT_SECONDS,
            )
        except subprocess.TimeoutExpired as exc:
            process.kill()
            process.communicate()
            raise ExportFailure(
                "timeout", "Document export took too long. Try a shorter draft.") from exc
        if process.returncode:
            try:
                failure = json.loads(error)
                raise ExportFailure(failure["code"], failure["message"])
            except (ValueError, KeyError, TypeError):
                raise ExportFailure(
                    "render_failed", "The approved document could not be exported. Please retry.") from None
        if not output or len(output) > MAX_EXPORT_BYTES:
            raise ExportFailure(
                "limit", "The document could not be exported within the download size limit.")
        return output
    finally:
        _WORKERS.release()


def _worker() -> None:
    try:
        # Even direct invocation of this worker cannot read unlimited input.
        raw = sys.stdin.buffer.read(1_000_001)
        if len(raw) > 1_000_000:
            raise ExportFailure(
                "limit", "The document exceeds the export input limit.")
        payload = json.loads(raw)
        result = _render(payload["document"], payload["format"])
        sys.stdout.buffer.write(result)
    except Exception as exc:
        failure = exc if isinstance(exc, ExportFailure) else ExportFailure(
            "render_failed", "The approved document could not be exported. Please retry."
        )
        sys.stderr.write(json.dumps(
            {"code": failure.code, "message": failure.message}))
        sys.exit(1)


if __name__ == "__main__":
    _worker()
