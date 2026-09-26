from copy import deepcopy
from io import BytesIO
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pytest

from app.services import pack_export


def representative_document():
    return {
        "audit": "INTERNAL-AUDIT-SECRET",
        "blocks": [
            {"id": "internal-id", "kind": "heading", "text": "Zoé Martin", "evidence": ["INTERNAL-EVIDENCE-SECRET"]},
            {"kind": "paragraph", "text": "zoe@example.test | +1 555 0100 | Paris"},
            {"kind": "heading", "text": "Selected experience"},
            {"kind": "paragraph", "text": "Software engineer — Example Cooperative\n2022–2025"},
            {"kind": "bullet", "text": "Built Python services with accessible review workflows.", "origin": "INTERNAL-ORIGIN-SECRET"},
            {"kind": "heading", "text": "Projects"},
            {"kind": "paragraph", "text": "Community library: searchable lending records & café catalogue."},
            {"kind": "paragraph", "text": '<img src="https://example.invalid/secret"> remains literal text.'},
        ],
    }


def test_pdf_export_selectable_exact_text_and_no_evidence_or_active_content():
    source = representative_document()
    original = deepcopy(source)
    result = pack_export.render_document(source, "pdf")
    assert result.startswith(b"%PDF-")
    reader = PdfReader(BytesIO(result))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    for value in ("Zoé Martin", "zoe@example.test", "2022–2025", "café catalogue", '<img src="https://example.invalid/secret">'):
        assert value in text
    assert "INTERNAL-" not in text
    assert b"INTERNAL-" not in result
    assert reader.pages[0].get("/Annots") is None
    assert "/OpenAction" not in reader.trailer["/Root"]
    assert reader.metadata.author == ""
    assert source == original
    assert float(reader.pages[0].mediabox.width) == pytest.approx(595.276, abs=0.01)


def test_docx_export_exact_unicode_headings_pagination_and_private_metadata():
    source = representative_document()
    source["blocks"].append({"kind": "paragraph", "text": "Unicode retained: مرحبا 世界"})
    result = pack_export.render_document(source, "docx")
    document = Document(BytesIO(result))
    assert [p.text for p in document.paragraphs] == [b["text"] for b in source["blocks"]]
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[2].style.name == "Heading 1"
    assert document.styles["Title"].paragraph_format.keep_with_next is True
    assert document.styles["Heading 1"].paragraph_format.keep_with_next is True
    assert all(p.paragraph_format.widow_control for p in document.paragraphs)
    assert document.paragraphs[4].style.name == "List Bullet"
    assert document.sections[0].page_width.mm == pytest.approx(210, abs=0.1)
    assert document.sections[0].page_height.mm == pytest.approx(297, abs=0.1)
    with ZipFile(BytesIO(result)) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            assert b"INTERNAL-" not in data
            if name.endswith(".rels"):
                assert b'TargetMode="External"' not in data
        assert b'PAGE' in archive.read("word/footer1.xml")
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""


def test_legacy_bullet_markers_render_once_with_subheading_in_both_formats():
    source = {"blocks": [
        {"kind": "heading", "text": "Curriculum vitae"},
        {"kind": "heading", "text": "Experience"},
        {"kind": "subheading", "text": "Backend Engineer — Cedar Demo"},
        {"kind": "bullet", "text": "•• Developed Python services."},
    ]}
    pdf = PdfReader(BytesIO(pack_export.render_document(source, "pdf")))
    pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Developed Python services." in pdf_text
    assert "••" not in pdf_text
    docx = Document(BytesIO(pack_export.render_document(source, "docx")))
    assert docx.paragraphs[2].style.name == "Heading 2"
    assert docx.paragraphs[3].style.name == "List Bullet"
    assert docx.paragraphs[3].text == "Developed Python services."
    assert source["blocks"][3]["text"] == "•• Developed Python services."


def test_pdf_multipage_content_stays_in_order_and_headings_stay_with_text():
    blocks = []
    for number in range(1, 35):
        blocks.extend([
            {"kind": "heading", "text": f"Project {number:02d}"},
            {"kind": "paragraph", "text": f"Evidence of project {number:02d}. " + "Accessible software and readable documentation. " * 9},
        ])
    result = pack_export._render({"blocks": blocks}, "pdf")
    reader = PdfReader(BytesIO(result))
    assert 2 <= len(reader.pages) <= pack_export.MAX_PAGES
    all_text = "\n".join(page.extract_text() for page in reader.pages)
    positions = [all_text.index(f"Project {number:02d}") for number in range(1, 35)]
    assert positions == sorted(positions)
    for page in reader.pages:
        text = page.extract_text()
        for number in range(1, 35):
            if f"Project {number:02d}" in text:
                assert f"Evidence of project {number:02d}." in text


@pytest.mark.parametrize("document", [
    {"blocks": []},
    {"blocks": [{"kind": "paragraph", "text": "x"}] * 101},
    {"blocks": [{"kind": "paragraph", "text": "x" * 2001}]},
    {"blocks": [{"kind": "paragraph", "text": "x" * 2000}] * 16},
    {"blocks": [{"kind": "image", "text": "https://example.invalid"}]},
    {"blocks": [{"kind": "paragraph", "text": "bad\x00text"}]},
])
def test_export_rejects_invalid_or_unbounded_input(document):
    with pytest.raises(pack_export.ExportFailure):
        pack_export.render_document(document, "pdf")


def test_export_glyph_failure_is_explicit_and_does_not_substitute_text():
    with pytest.raises(pack_export.ExportFailure, match="Download DOCX") as error:
        pack_export.render_document({"blocks": [{"kind": "paragraph", "text": "世界"}]}, "pdf")
    assert error.value.code == "unsupported_glyph"


def test_export_page_and_output_limits(monkeypatch):
    monkeypatch.setattr(pack_export, "MAX_PAGES", 1)
    blocks = [{"kind": "paragraph", "text": "Readable text. " * 100}] * 10
    with pytest.raises(pack_export.ExportFailure, match="pages"):
        pack_export._render({"blocks": blocks}, "pdf")
    monkeypatch.setattr(pack_export, "MAX_EXPORT_BYTES", 10)
    with pytest.raises(pack_export.ExportFailure, match="5 MB"):
        pack_export._BoundedBuffer().write(b"x" * 11)


def test_export_worker_timeout_kills_and_reaps_process(monkeypatch):
    class TimedOutProcess:
        calls = 0
        killed = False

        def communicate(self, data=None, timeout=None):
            self.calls += 1
            if self.calls == 1:
                assert timeout == pack_export.EXPORT_TIMEOUT_SECONDS
                raise subprocess.TimeoutExpired("renderer", timeout)
            return b"", b""

        def kill(self):
            self.killed = True

    process = TimedOutProcess()
    monkeypatch.setattr(pack_export.subprocess, "Popen", lambda *args, **kwargs: process)
    with pytest.raises(pack_export.ExportFailure, match="too long") as error:
        pack_export.render_document(representative_document(), "pdf")
    assert error.value.code == "timeout"
    assert process.killed and process.calls == 2


def test_export_deadline_and_unknown_format():
    with pytest.raises(pack_export.ExportFailure) as error:
        pack_export._check_deadline(0)
    assert error.value.code == "timeout"
    with pytest.raises(pack_export.ExportFailure) as error:
        pack_export.render_document(representative_document(), "html")
    assert error.value.code == "unsupported"
