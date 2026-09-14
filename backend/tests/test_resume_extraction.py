import io
import time
import uuid

import pytest
from docx import Document
from fastapi import status
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import select

from app.core.config import get_settings
from app.core.db import get_db
from app.core.security import create_access_token, hash_password
from app.main import create_application
from app.models import ResumeExtraction, User
from app.services import resume_extraction, resume_store


def pdf_bytes(*page_texts: str) -> bytes:
    objects: list[bytes] = []
    page_refs = " ".join(f"{3 + index * 2} 0 R" for index in range(len(page_texts)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{page_refs}] /Count {len(page_texts)} >>".encode())
    font_id = 3 + len(page_texts) * 2
    for index, text in enumerate(page_texts):
        content_id = 4 + index * 2
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def scanned_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)
    return buffer.getvalue()


def encrypted_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Ada Lovelace")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Years"
    table.cell(1, 0).text = "Engineer"
    table.cell(1, 1).text = "5"
    document.add_paragraph("Python and systems design")
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def extraction_client(db_session):
    app = create_application()

    def override():
        yield db_session

    app.dependency_overrides[get_db] = override
    return TestClient(app, raise_server_exceptions=False)


@pytest.fixture()
def extraction_users(db_session):
    users = [
        User(email=f"extract-{name}@jobpilot-test.com", password_hash=hash_password("password-123"))
        for name in ("owner", "other")
    ]
    db_session.add_all(users)
    db_session.commit()
    for user in users:
        db_session.refresh(user)
    return users


@pytest.fixture()
def extraction_storage(tmp_path, monkeypatch):
    target = tmp_path / "resumes"
    monkeypatch.setattr(get_settings(), "RESUME_STORAGE_DIR", str(target))
    return target


def headers(user: User) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(user.id)}"}


def upload(client, user, filename, data):
    return client.post(
        "/api/resumes",
        headers=headers(user),
        files={"file": (filename, data, "application/octet-stream")},
    )


def test_extracts_pdf_pages_and_preserves_original_bytes(
    extraction_client, extraction_users, extraction_storage
):
    owner, _ = extraction_users
    source = pdf_bytes("First page", "Second page")
    uploaded = upload(extraction_client, owner, "cv.pdf", source)
    assert uploaded.status_code == 201
    resume_id = uploaded.json()["id"]

    response = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner))
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "succeeded"
    assert "--- Page 1 ---" in body["original_text"]
    assert "First page" in body["original_text"]
    assert "--- Page 2 ---" in body["original_text"]
    assert "Second page" in body["original_text"]
    assert body["parser_name"] == "pypdf"
    assert extraction_storage.joinpath(resume_id).read_bytes() == source


def test_extracts_docx_paragraphs_and_tables_in_order(
    extraction_client, extraction_users, extraction_storage
):
    owner, _ = extraction_users
    source = docx_bytes()
    resume_id = upload(extraction_client, owner, "cv.docx", source).json()["id"]
    body = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    text = body["original_text"]
    assert text.index("Ada Lovelace") < text.index("Role\tYears") < text.index("Python and systems design")
    assert "Engineer\t5" in text
    assert extraction_storage.joinpath(resume_id).read_bytes() == source


@pytest.mark.parametrize(
    ("source", "failure_code"),
    [(scanned_pdf_bytes(), "ocr_required"), (encrypted_pdf_bytes(), "encrypted")],
)
def test_pdf_failures_are_honest(extraction_client, extraction_users, source, failure_code):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", source).json()["id"]
    body = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert body["status"] == "failed"
    assert body["failure_code"] == failure_code
    assert body["original_text"] is None


def test_processing_limits_and_failed_retry(extraction_client, extraction_users, monkeypatch):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("One", "Two")).json()["id"]
    monkeypatch.setattr(get_settings(), "RESUME_EXTRACTION_MAX_PAGES", 1)
    failed = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert failed["failure_code"] == "too_complex"

    monkeypatch.setattr(get_settings(), "RESUME_EXTRACTION_MAX_PAGES", 2)
    retried = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert retried["status"] == "succeeded"


def test_output_limit_is_persisted_as_failure(extraction_client, extraction_users, monkeypatch):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("Long resume text")).json()["id"]
    monkeypatch.setattr(get_settings(), "RESUME_EXTRACTION_MAX_CHARS", 10)
    body = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert body["status"] == "failed"
    assert body["failure_code"] == "output_too_large"


def test_hard_timeout_returns_without_waiting_for_parser(monkeypatch):
    def slow_parser(*args, **kwargs):
        time.sleep(0.2)
        return "late", "parser", "1"

    monkeypatch.setattr(resume_extraction, "extract_pdf", slow_parser)
    started = time.monotonic()
    with pytest.raises(resume_extraction.ExtractionFailure) as error:
        resume_extraction.extract(
            pdf_bytes("Slow"),
            "pdf",
            timeout_seconds=0.02,
            max_chars=100,
            max_pages=1,
            max_blocks=1,
        )
    assert error.value.code == "timeout"
    assert time.monotonic() - started < 0.15


def test_timeout_failure_can_retry(extraction_client, extraction_users, monkeypatch):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("Retry me")).json()["id"]
    real_extract = resume_extraction.extract
    monkeypatch.setattr(
        resume_extraction,
        "extract",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            resume_extraction.ExtractionFailure("timeout", "Text extraction took too long.")
        ),
    )
    assert extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()["failure_code"] == "timeout"
    monkeypatch.setattr(resume_extraction, "extract", real_extract)
    assert extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()["status"] == "succeeded"


def test_edit_confirm_and_repeat_extract_do_not_overwrite_review(
    extraction_client, extraction_users
):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("Original")).json()["id"]
    extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner))
    edited = extraction_client.patch(
        f"/api/resumes/{resume_id}/extraction",
        headers=headers(owner),
        json={"draft_text": "Corrected CV text"},
    ).json()
    assert edited["original_text"] != edited["draft_text"]
    assert edited["reviewed_at"] is None
    confirmed = extraction_client.post(
        f"/api/resumes/{resume_id}/extraction/confirm", headers=headers(owner)
    ).json()
    assert confirmed["reviewed_at"] is not None
    persisted = extraction_client.get(
        f"/api/resumes/{resume_id}/extraction", headers=headers(owner)
    ).json()
    assert persisted["draft_text"] == "Corrected CV text"
    assert persisted["reviewed_at"] == confirmed["reviewed_at"]

    repeated = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert repeated["draft_text"] == "Corrected CV text"
    assert repeated["reviewed_at"] == confirmed["reviewed_at"]
    changed = extraction_client.patch(
        f"/api/resumes/{resume_id}/extraction",
        headers=headers(owner),
        json={"draft_text": "Changed again"},
    ).json()
    assert changed["reviewed_at"] is None


def test_corrupt_stored_source_has_useful_failure(
    extraction_client, extraction_users, extraction_storage
):
    owner, _ = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("Initially valid")).json()["id"]
    extraction_storage.joinpath(resume_id).write_bytes(b"%PDF-corrupt\n%%EOF")
    body = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    assert body["status"] == "failed"
    assert body["failure_code"] == "malformed"
    assert body["failure_message"]


def test_unsupported_parser_format_is_rejected():
    with pytest.raises(resume_extraction.ExtractionFailure) as error:
        resume_extraction.extract(
            b"data",
            "rtf",
            timeout_seconds=1,
            max_chars=100,
            max_pages=1,
            max_blocks=1,
        )
    assert error.value.code == "unsupported"


def test_extraction_ownership_and_resume_delete_cascade(
    extraction_client, extraction_users, db_session, extraction_storage
):
    owner, other = extraction_users
    resume_id = upload(extraction_client, owner, "cv.pdf", pdf_bytes("Private")).json()["id"]
    extraction = extraction_client.post(f"/api/resumes/{resume_id}/extract", headers=headers(owner)).json()
    for method, path in [
        ("get", f"/api/resumes/{resume_id}/extraction"),
        ("post", f"/api/resumes/{resume_id}/extract"),
        ("patch", f"/api/resumes/{resume_id}/extraction"),
        ("post", f"/api/resumes/{resume_id}/extraction/confirm"),
    ]:
        response = getattr(extraction_client, method)(
            path,
            headers=headers(other),
            **({"json": {"draft_text": "stolen"}} if method == "patch" else {}),
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND
        assert response.json()["detail"] == "Resume not found"

    assert extraction_client.delete(f"/api/resumes/{resume_id}", headers=headers(owner)).status_code == 204
    assert db_session.scalar(
        select(ResumeExtraction).where(ResumeExtraction.id == uuid.UUID(extraction["id"]))
    ) is None
    assert not extraction_storage.joinpath(resume_id).exists()
